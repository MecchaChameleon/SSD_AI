from __future__ import annotations

import csv
import json
import os
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", str(Path(__file__).resolve().parent / ".matplotlib"))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
SUMMARY = ROOT / "summary"
RAW = ROOT / "raw"
DOCS = ROOT / "docs"
CHARTS = DOCS / "charts"


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def latest_load_result(folder: Path, filename: str) -> dict:
    candidates = sorted(folder.glob(f"*/{filename}"))
    return json.loads(candidates[-1].read_text(encoding="utf-8")) if candidates else {}


def save_chart(fig, name: str) -> Path:
    CHARTS.mkdir(parents=True, exist_ok=True)
    path = CHARTS / name
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_charts(rows: list[dict], latency: list[dict], load: dict) -> list[Path]:
    plt.rcParams.update({"font.family": "Malgun Gothic", "axes.unicode_minus": False})
    models = {"A_cases_as_annual": "거래건수=연간", "B_cases_as_monthly": "거래건수=월간"}
    fig, ax = plt.subplots(figsize=(12, 6.75))
    for key, label in models.items():
        selected = [r for r in rows if r["case_model"] == key]
        ax.plot([int(r["year"]) for r in selected], [int(float(r["backend_requests_year"])) for r in selected], marker="o", linewidth=2.5, label=label)
    ax.set(title="연도별 Backend 요청량", xlabel="사업연차", ylabel="연간 요청 수")
    ax.grid(alpha=.25); ax.legend(); fig.tight_layout()
    p1 = save_chart(fig, "01_backend_requests_16x9.png")

    fig, ax = plt.subplots(figsize=(12, 6.75))
    selected = [r for r in rows if r["case_model"] == "B_cases_as_monthly"]
    x = [int(r["year"]) for r in selected]
    ax.bar([v-.18 for v in x], [float(r["db_total_queries_year"])/1e6 for r in selected], .36, label="DB query (백만 회)")
    ax.bar([v+.18 for v in x], [float(r["s3_get_year"])/1e6 for r in selected], .36, label="S3 GET 환산 (백만 회)")
    ax.set(title="월 거래건수 해석 시 연간 DB·S3 사용량", xlabel="사업연차", ylabel="백만 회")
    ax.grid(axis="y", alpha=.25); ax.legend(); fig.tight_layout()
    p2 = save_chart(fig, "02_db_s3_usage_16x9.png")

    fig, ax = plt.subplots(figsize=(12, 6.75))
    names = [r["endpoint_name"] for r in latency]
    vals = [float(r["p95_latency_ms"]) for r in latency]
    ax.barh(names, vals, color="#2f7ed8")
    ax.set(title="읽기 전용 API 응답시간 (각 20회, p95)", xlabel="milliseconds")
    ax.grid(axis="x", alpha=.25); fig.tight_layout()
    p3 = save_chart(fig, "03_read_api_p95_16x9.png")

    stages = load.get("stages", [])
    fig, ax1 = plt.subplots(figsize=(12, 6.75))
    x = [s["concurrent_users"] for s in stages]
    ax1.plot(x, [s["throughput_rps"] for s in stages], marker="o", color="#2f7ed8", label="RPS")
    ax1.set(xlabel="동시 사용자", ylabel="RPS", title="기존 혼합 부하 실험 (구매 5%)")
    ax2 = ax1.twinx(); ax2.plot(x, [s["latency_p95_ms"] for s in stages], marker="s", color="#ef6c35", label="p95")
    ax2.axhline(3000, color="#888", linestyle="--", linewidth=1); ax2.set_ylabel("p95 ms")
    ax1.grid(alpha=.25); fig.tight_layout()
    p4 = save_chart(fig, "04_historical_load_16x9.png")
    return [p1, p2, p3, p4]


def fmt(n: str | float, digits=0) -> str:
    return f"{float(n):,.{digits}f}"


def add_table(doc: Document, headers: list[str], data: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
    for i, h in enumerate(headers): table.rows[0].cells[i].text = h
    for row in data:
        cells = table.add_row().cells
        for i, value in enumerate(row): cells[i].text = str(value)


def main() -> None:
    rows = read_csv(SUMMARY / "infrastructure_summary.csv")
    latency = read_csv(RAW / "backend_latency.csv")
    load = latest_load_result(ROOT.parent / "simulation/load-results-users", "concurrent_users.json")
    charts = make_charts(rows, latency, load)
    measured_at = json.loads((RAW / "measurement.json").read_text(encoding="utf-8"))["measured_at_utc"]

    md = f"""# 실제 Backend 기반 인프라 사용량 측정·환산 보고서

측정시각: {measured_at}  
대상: `https://ssd-be.onrender.com`  

## 1. 결론

- 읽기 전용 실제 API 10종을 각각 20회, 총 200회 호출했으며 성공률은 100%였다. 구매 생성과 GPT API 호출은 수행하지 않았다.
- 요청 구성은 일반 조회 98.99%, 구매 1.00%, GPT 기능 0.01%로 두었다. 거래건수가 구매 요청 수이므로 전체 요청은 `거래건수 / 0.01`, 즉 거래 1건당 요청 100건으로 역산했다.
- 거래건수 900/3,600/10,800의 기간이 불명확하여, 연간 수치로 보는 A안과 월간 수치로 보는 B안을 모두 계산했다.
- 더 보수적인 B안 3년차도 연간 Backend 요청 12,960,000건, 정상 활동시간 평균 0.62 RPS, 피크 가정 6.16 RPS다. 기존 혼합 부하 실험의 처리량 약 35 RPS보다 낮다.
- 따라서 현재 사업계획 범위에서는 종량 사용량보다 Render·DB의 고정 최소요금, 무료 DB 만료/용량·연결 제한, 이미지 저장정책이 비용과 안정성에 더 큰 영향을 줄 가능성이 높다.

## 2. 실험 설계

- A(직접 측정): 실제 배포 API의 상태코드, 응답 바이트, 지연시간. GET만 실행.
- B(측정 기반 환산): 실측 응답 크기와 코드상 DB query 추적치를 사업연도 요청량에 곱함.
- C(사업 가정): 요청 비율, 피크 10배, 상품/매장/연 12개, 이미지 2장/상품, 이미지 500KB, GPT 출력 500 token.
- D(측정 불가): Render CPU/RAM, DB 실제 크기·최대 connection·`pg_stat_statements`. 관리 콘솔/DB 권한 없이는 확정할 수 없다.
- GPT는 코드의 프롬프트 구조와 호출 빈도만 계산했다. 외부 GPT API 실제 호출은 0회다.

## 3. 핵심 수치

| 해석 | 연차 | 연간 거래 | 연간 Backend 요청 | 피크 RPS | 연간 DB query | GPT 기능 카운트 |
|---|---:|---:|---:|---:|---:|---:|
"""
    for r in rows:
        md += f"| {'연간' if r['case_model'].startswith('A_') else '월간→연환산'} | {r['year']} | {fmt(r['cases_per_year'])} | {fmt(r['backend_requests_year'])} | {fmt(r['peak_rps'],2)} | {fmt(r['db_total_queries_year'])} | {fmt(r['gpt_simulated_calls_year'])} |\n"
    md += """

## 4. 해석상 주의

- 실제 API 응답에서 URL은 관측됐지만 현재 표본은 S3/CloudFront URL이 아니었다. 그러므로 S3 GET 실측 0을 ‘서비스의 S3 사용량 0’으로 해석하면 안 된다. 보고서의 S3 GET은 응답 내 자산 URL 수를 이용한 B 환산이고, PUT·저장량은 C 가정이다.
- 현재 두 환경변수에 동일한 SELLER JWT가 들어갔지만 배포 서버가 조사 대상 GET 모두에 200을 반환했다. 역할별 인가가 의도된 것이라면 별도 보안 검증이 필요하다.
- 기존 혼합 부하실험은 구매 5%였고 과거 실제 구매 데이터가 생성된 실험이다. 이번 사용량 측정에서는 재실행하지 않고 역사적 결과만 인용했다.
- 비용 산출 시 사용량 표와 가격표를 분리해야 한다. 공급자·리전·플랜과 무료 DB의 만료 조건을 확정한 뒤 `월 사용량 × 단가 + 고정요금`으로 계산한다.

## 5. 산출물

- 원시 HTTP: `raw/backend_requests.csv`, `raw/backend_latency.csv`
- DB/S3/GPT: `raw/db_queries.csv`, `raw/s3_requests.csv`, `raw/gpt_simulation.csv`
- 연도별 결과: `summary/*_year*_summary.csv`
- 통합표: `summary/infrastructure_summary.csv`
- PPT용 16:9 PNG: `docs/charts/*.png`
"""
    DOCS.mkdir(exist_ok=True)
    (DOCS / "infrastructure_measurement_report.md").write_text(md, encoding="utf-8")

    doc = Document(); sec = doc.sections[0]; sec.top_margin = Inches(.6); sec.bottom_margin = Inches(.6)
    title = doc.add_heading("실제 Backend 기반 인프라 사용량 측정·환산 보고서", 0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"측정시각 {measured_at}\n대상 https://ssd-be.onrender.com"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("1. 실험 결론", 1)
    for text in [
        "실제 읽기 API 10종 × 20회(총 200회)는 모두 성공했다. 이번 측정에서 구매 생성 및 GPT 외부 호출은 0회다.",
        "거래 1건이 전체 요청의 1%라는 가정이므로 거래 1건당 전체 요청 100건으로 환산했다.",
        "보수적인 월간 거래건수 해석(B안) 3년차 피크도 6.16 RPS로, 과거 혼합 부하 실험 처리량 약 35 RPS보다 낮다.",
        "현재 규모에서는 트래픽 종량비보다 서버·DB 고정요금과 무료 DB 제약이 더 중요한 비용 요인이다.",
    ]: doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("2. 측정 신뢰도 구분", 1)
    add_table(doc, ["등급", "의미"], [["A", "실제 배포 API 직접 측정"], ["B", "실측값 기반 환산"], ["C", "사업·시뮬레이션 가정"], ["D", "권한 부족으로 측정 불가"]])
    doc.add_heading("3. 연도별 통합 결과", 1)
    add_table(doc, ["모델", "연차", "연 거래", "Backend 요청", "피크 RPS", "DB query", "GPT 카운트"], [["연간" if r['case_model'].startswith('A_') else "월간→연환산", r['year'], fmt(r['cases_per_year']), fmt(r['backend_requests_year']), fmt(r['peak_rps'],2), fmt(r['db_total_queries_year']), fmt(r['gpt_simulated_calls_year'])] for r in rows])
    for i, chart in enumerate(charts, 1):
        doc.add_page_break(); doc.add_heading(f"그림 {i}", 1); doc.add_picture(str(chart), width=Inches(7.1))
    doc.add_page_break(); doc.add_heading("4. 한계와 다음 조치", 1)
    for text in [
        "Render CPU/RAM, DB 크기·연결 수는 D(측정 불가)다. Render Metrics와 PostgreSQL 통계를 확보해야 한다.",
        "S3 실측 표본이 없어 PUT·저장량은 C 가정이다. 실제 업로드 로그 또는 버킷 지표로 교체한다.",
        "공급자·리전·플랜을 확정한 뒤 통합 CSV에 단가를 곱해 비용을 계산한다. 무료 티어도 비용 0이 아니라 한도와 만료 위험을 별도 표기한다.",
        "동일 SELLER JWT로 구매자 GET도 허용됐다. 의도된 정책이 아니라면 인가 테스트가 필요하다.",
    ]: doc.add_paragraph(text, style="List Bullet")
    styles = doc.styles
    styles["Normal"].font.name = "Malgun Gothic"; styles["Normal"].font.size = Pt(9)
    out = DOCS / "인프라_사용량_측정_및_3개년_환산_보고서.docx"; doc.save(out)
    print(out)


if __name__ == "__main__":
    main()

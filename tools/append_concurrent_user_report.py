from __future__ import annotations

import json
import os
from pathlib import Path

os.environ.setdefault("MPLBACKEND", "Agg")
os.environ.setdefault("MPLCONFIGDIR", str(Path(__file__).resolve().parents[1] / ".runtime" / "matplotlib"))

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "simulation" / "배포_서버_구매_API_부하테스트_결과.docx"
READ_FILES = [
    ROOT / "simulation/load-results-users/20260831T144706/concurrent_users.json",
    ROOT / "simulation/load-results-users/20260831T144801/concurrent_users.json",
]
MIXED_FILES = [
    ROOT / "simulation/load-results-users/20260831T145022/concurrent_users.json",
    ROOT / "simulation/load-results-users/20260831T145221/concurrent_users.json",
]
CHART_DIR = ROOT / "simulation/load-results-users/report-charts"


def rows(paths: list[Path]) -> list[dict]:
    values = []
    for path in paths:
        values.extend(json.loads(path.read_text(encoding="utf-8"))["stages"])
    return sorted(values, key=lambda row: row["concurrent_users"])


def add_table(document: Document, data: list[dict]) -> None:
    headers = ["동시 사용자", "요청", "성공률", "RPS", "평균(ms)", "p50(ms)", "p95(ms)", "p99(ms)"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for row in data:
        values = [row["concurrent_users"], row["requests"], f"{row['success_rate']*100:.1f}%",
                  f"{row['throughput_rps']:.2f}", f"{row['latency_avg_ms']:.2f}",
                  f"{row['latency_p50_ms']:.2f}", f"{row['latency_p95_ms']:.2f}", f"{row['latency_p99_ms']:.2f}"]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = str(value)


def chart(data: list[dict], title: str, filename: str) -> Path:
    users = [row["concurrent_users"] for row in data]
    fig, ax1 = plt.subplots(figsize=(10, 5.4))
    ax1.plot(users, [row["throughput_rps"] for row in data], marker="o", color="#2563EB", linewidth=2.3)
    ax1.set_xlabel("동시 사용자 수")
    ax1.set_ylabel("처리량(RPS)", color="#2563EB")
    ax1.tick_params(axis="y", labelcolor="#2563EB")
    ax1.grid(alpha=.2)
    ax2 = ax1.twinx()
    ax2.plot(users, [row["latency_p95_ms"] for row in data], marker="s", color="#DC2626", linewidth=2.3)
    ax2.axhline(3000, color="#DC2626", linestyle="--", alpha=.55)
    ax2.set_ylabel("p95 응답시간(ms)", color="#DC2626")
    ax2.tick_params(axis="y", labelcolor="#DC2626")
    plt.title(title)
    fig.tight_layout()
    path = CHART_DIR / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return path


def main() -> None:
    plt.rcParams["font.family"] = ["Malgun Gothic", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    read = rows(READ_FILES)
    mixed = rows(MIXED_FILES)
    read_chart = chart(read, "조회 중심 동시접속 테스트", "concurrent_users_read_only.png")
    mixed_chart = chart(mixed, "조회 + 실제 구매 5% 혼합 테스트", "concurrent_users_mixed.png")

    document = Document(DOCX)
    document.add_page_break()
    document.add_heading("6. 동시접속 사용자 테스트", level=1)
    document.add_paragraph(
        "단순 동시 요청 수가 아니라 가상 사용자가 0.5~1.5초의 체류시간을 두고 상품 목록, 상세, 지도, 구매 내역, 판매자 상품 목록을 반복 조회하도록 구성했다. "
        "각 단계는 약 20초간 유지했으며 성공률 99% 미만 또는 전체 p95 3초 초과 시 중단했다."
    )

    document.add_heading("6.1 조회 중심 시나리오", level=2)
    add_table(document, read)
    document.add_picture(str(read_chart), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "100명에서 39.67 RPS, 성공률 100%, p95 2.88초로 중단 기준 이내였다. "
        "125명에서는 처리량이 36.42 RPS로 감소하고 p95가 5.35초로 증가했다. 따라서 조회 중심 포화 구간은 100~125명 사이이며, p95 1초 목표의 권장 동시접속자는 30명 이하, p95 3초 허용 시 최대 약 100명이다."
    )

    document.add_heading("6.2 조회 + 실제 구매 5% 혼합 시나리오", level=2)
    add_table(document, mixed)
    document.add_picture(str(mixed_chart), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    purchase_count = sum(row.get("endpoints", {}).get("purchase", {}).get("requests", 0) for row in mixed)
    document.add_paragraph(
        f"혼합 테스트에서 실제 구매 {purchase_count}건이 생성됐다. 50명에서는 35.15 RPS, 전체 p95 0.83초로 안정적이었다. "
        "75명에서는 전체 p95 2.81초지만 구매 API 자체 p95가 5.66초로 느려졌고, 100명에서는 전체 p95 4.26초로 중단 기준을 초과했다. "
        "따라서 실제 운영형 권장 동시접속자는 50명 이하이며 75명은 순간 피크 허용 범위, 100명은 현재 무료 구성의 비권장 영역이다."
    )

    document.add_heading("6.3 종합 용량 판단", level=2)
    for value in [
        "조회 중심: 권장 30명 이하, 조건부 허용 100명, 포화 시작 125명.",
        "구매 5% 혼합: 권장 50명 이하, 조건부 피크 75명, 100명에서 응답시간 기준 실패.",
        "순수 동시 결제: p95 1초 기준 동시성 3, p95 3초 기준 동시성 5.",
        "HTTP 오류는 관측되지 않았지만 포화 구간에서는 처리량이 정체되고 대기시간이 급증했다.",
        "무료 DB와 무료 인스턴스 결과이므로 실제 출시 용량은 운영 등급 인프라에서 10~30분 지속 테스트로 다시 확인해야 한다.",
    ]:
        document.add_paragraph(value, style="List Bullet")
    document.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()

from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("MPLBACKEND", "Agg")
os.environ.setdefault("MPLCONFIGDIR", str(Path(__file__).resolve().parents[1] / ".runtime" / "matplotlib"))

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from append_concurrent_user_report import CHART_DIR, MIXED_FILES, READ_FILES, add_table, chart, rows


def main() -> None:
    plt.rcParams["font.family"] = ["Malgun Gothic", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    read = rows(READ_FILES)
    mixed = rows(MIXED_FILES)
    read_chart = chart(read, "조회 중심 동시접속 테스트", "concurrent_users_read_only.png")
    mixed_chart = chart(mixed, "조회 + 실제 구매 5% 혼합 테스트", "concurrent_users_mixed.png")
    output = CHART_DIR.parents[1] / "동시접속_부하테스트_결과_최종본.docx"

    document = Document()
    document.styles["Normal"].font.name = "맑은 고딕"
    document.styles["Normal"].font.size = Pt(10)
    title = document.add_heading("동시접속 사용자 부하 테스트 결과", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary = document.add_paragraph("조회 중심 최대 125명 · 구매 5% 혼합 최대 100명 측정")
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("1. 핵심 결론", level=1)
    for value in [
        "조회 중심 권장 동시접속: 30명 이하",
        "조회 중심 조건부 허용: 최대 100명, 125명에서 p95 5.35초로 포화",
        "구매 5% 혼합 권장 동시접속: 50명 이하",
        "구매 혼합 75명은 순간 피크 범위, 100명에서 p95 4.26초로 기준 초과",
        "전체 동시접속 테스트 HTTP 성공률: 100%",
    ]:
        document.add_paragraph(value, style="List Bullet")

    document.add_heading("2. 조회 중심 동시접속 결과", level=1)
    add_table(document, read)
    document.add_picture(str(read_chart), width=Inches(6.4))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "100명에서 39.67 RPS, p95 2.88초였으며 125명에서는 처리량이 36.42 RPS로 감소하고 p95가 5.35초로 증가했다. "
        "따라서 조회 중심 포화 구간은 100~125명 사이다."
    )

    document.add_heading("3. 조회 + 실제 구매 5% 혼합 결과", level=1)
    add_table(document, mixed)
    document.add_picture(str(mixed_chart), width=Inches(6.4))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    purchases = sum(row.get("endpoints", {}).get("purchase", {}).get("requests", 0) for row in mixed)
    document.add_paragraph(
        f"총 실제 구매 {purchases}건이 포함됐다. 50명에서 35.15 RPS, p95 0.83초로 안정적이었다. "
        "75명은 전체 p95 2.81초, 100명은 p95 4.26초로 자동 중단 기준을 초과했다."
    )

    document.add_heading("4. 테스트 조건과 한계", level=1)
    for value in [
        "각 단계 약 20초 유지, 사용자별 요청 간격 0.5~1.5초",
        "상품 목록·상세·지도·구매 내역·판매자 상품 목록을 실제 사용자 비율로 혼합",
        "중단 기준: 성공률 99% 미만 또는 전체 p95 3초 초과",
        "Render 무료 인스턴스와 무료 DB 결과이므로 운영 등급 인프라에서는 지속 부하 재측정 필요",
    ]:
        document.add_paragraph(value, style="List Bullet")

    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
